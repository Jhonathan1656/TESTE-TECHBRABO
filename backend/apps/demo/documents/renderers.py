import io
import zipfile
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal

import openpyxl
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

_FIXED_ZIP_TIMESTAMP = (2026, 1, 1, 0, 0, 0)
_HEADER_BACKGROUND = colors.HexColor("#E8EEF7")
_CURRENCY_FORMAT = '"R$" #,##0.00'


@dataclass(frozen=True)
class GridColumn:
    label: str
    values: tuple[str, ...]


@dataclass(frozen=True)
class GridColumnGroup:
    title: str
    columns: tuple[GridColumn, ...]


@dataclass(frozen=True)
class PriceGridDocument:
    heading: str
    info_lines: tuple[str, ...]
    band_header: str
    band_labels: tuple[str, ...]
    groups: tuple[GridColumnGroup, ...]
    footer: tuple[str, ...]


@dataclass(frozen=True)
class PriceRow:
    name: str
    accommodation: str
    values: tuple[Decimal, ...]


@dataclass(frozen=True)
class PriceRowsWorkbook:
    heading: str
    info_lines: tuple[str, ...]
    header: tuple[str, ...]
    rows: tuple[PriceRow, ...]
    footer: tuple[str, ...]
    sheet_title: str


@dataclass(frozen=True)
class NoticeSection:
    title: str
    lines: tuple[str, ...]


@dataclass(frozen=True)
class NoticeDocument:
    heading: str
    title: str
    paragraphs: tuple[str, ...]
    sections: tuple[NoticeSection, ...]
    footer: tuple[str, ...]


def render_price_grid_pdf(document: PriceGridDocument) -> bytes:
    styles = getSampleStyleSheet()
    story = [
        Paragraph(document.heading, styles["Heading2"]),
        *(Paragraph(line, styles["Normal"]) for line in document.info_lines),
        Spacer(1, 5 * mm),
        build_grid_table(document),
        Spacer(1, 4 * mm),
        *(Paragraph(line, styles["Italic"]) for line in document.footer),
    ]
    return _build_pdf(story, landscape(A4), document.heading)


def build_grid_table(document: PriceGridDocument) -> Table:
    columns = [column for group in document.groups for column in group.columns]
    header_rows, spans = _build_header(document)
    body = [
        [band_label, *(column.values[index] for column in columns)]
        for index, band_label in enumerate(document.band_labels)
    ]
    header_count = len(header_rows)
    table = Table(header_rows + body, repeatRows=header_count)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (-1, header_count - 1), _HEADER_BACKGROUND),
                ("FONTNAME", (0, 0), (-1, header_count - 1), "Helvetica-Bold"),
                ("ALIGN", (1, header_count), (-1, -1), "RIGHT"),
                ("ALIGN", (1, 0), (-1, header_count - 1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                *spans,
            ]
        )
    )
    return table


def render_price_rows_xlsx(workbook_data: PriceRowsWorkbook, created_at: datetime) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = workbook_data.sheet_title
    sheet.append([workbook_data.heading])
    sheet["A1"].font = Font(bold=True, size=13)
    for line in workbook_data.info_lines:
        sheet.append([line])
    sheet.append([])
    sheet.append(list(workbook_data.header))
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
    for row in workbook_data.rows:
        sheet.append([row.name, row.accommodation, *(float(value) for value in row.values)])
        for cell in sheet[sheet.max_row][2:]:
            cell.number_format = _CURRENCY_FORMAT
    sheet.append([])
    for line in workbook_data.footer:
        sheet.append([line])
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 16
    workbook.properties.created = created_at
    workbook.properties.modified = created_at
    buffer = io.BytesIO()
    workbook.save(buffer)
    return normalize_zip_timestamps(buffer.getvalue())


def render_notice_pdf(document: NoticeDocument) -> bytes:
    styles = getSampleStyleSheet()
    story = [
        Paragraph(document.heading, styles["Heading2"]),
        Paragraph(document.title, styles["Heading3"]),
        *(Paragraph(paragraph, styles["Normal"]) for paragraph in document.paragraphs),
        Spacer(1, 4 * mm),
    ]
    for section in document.sections:
        story.append(Paragraph(f"<b>{section.title}</b>", styles["Normal"]))
        story.extend(Paragraph(line, styles["Normal"]) for line in section.lines)
        story.append(Spacer(1, 3 * mm))
    story.extend(Paragraph(line, styles["Italic"]) for line in document.footer)
    return _build_pdf(story, A4, document.heading)


def normalize_zip_timestamps(data: bytes) -> bytes:
    """Rewrites a zip with fixed entry dates so identical content always yields identical bytes."""
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            normalized = zipfile.ZipInfo(entry.filename, date_time=_FIXED_ZIP_TIMESTAMP)
            normalized.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(normalized, source.read(entry.filename))
    return output.getvalue()


def _build_header(document: PriceGridDocument) -> tuple[list[list[str]], list[tuple]]:
    columns = [column for group in document.groups for column in group.columns]
    if not any(group.title for group in document.groups):
        return [[document.band_header, *(column.label for column in columns)]], []
    group_row = [document.band_header]
    label_row = [""]
    spans: list[tuple] = [("SPAN", (0, 0), (0, 1))]
    position = 1
    for group in document.groups:
        width = len(group.columns)
        group_row.extend([group.title, *[""] * (width - 1)])
        label_row.extend(column.label for column in group.columns)
        if width > 1:
            spans.append(("SPAN", (position, 0), (position + width - 1, 0)))
        position += width
    return [group_row, label_row], spans


def _build_pdf(story: list, page_size, title: str) -> bytes:
    buffer = io.BytesIO()
    pdf = SimpleDocTemplate(
        buffer,
        pagesize=page_size,
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=title,
        invariant=1,
    )
    pdf.build(story)
    return buffer.getvalue()
