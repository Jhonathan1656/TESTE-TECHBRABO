import csv
import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pdfplumber
from openpyxl.utils import get_column_letter

_CSV_SAMPLE_BYTES = 4096
_CSV_DELIMITER_PREFERENCE = (";", "\t")


@dataclass(frozen=True)
class TableCell:
    text: str
    number: Decimal | None = None
    reference: str = ""


@dataclass(frozen=True)
class RawTable:
    rows: tuple[tuple[TableCell, ...], ...]
    page: int | None = None
    sheet: str | None = None


@dataclass(frozen=True)
class DocumentContent:
    text: str
    tables: tuple[RawTable, ...]


class UnreadableDocumentError(Exception):
    pass


def read_document_content(path: Path, file_format: str) -> DocumentContent:
    return read_document_bytes(path.read_bytes(), file_format)


def read_document_bytes(data: bytes, file_format: str) -> DocumentContent:
    readers = {"pdf": read_pdf_content, "xlsx": read_xlsx_content, "csv": read_csv_content}
    try:
        return readers[file_format](data)
    except Exception as error:
        # Third-party parsers raise many unrelated types for corrupt or protected files.
        raise UnreadableDocumentError(str(error)) from error


def read_pdf_content(data: bytes) -> DocumentContent:
    page_texts: list[str] = []
    tables: list[RawTable] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            page_texts.append(page.extract_text() or "")
            tables.extend(
                build_raw_table(table, page=page_number) for table in page.extract_tables()
            )
    return DocumentContent(text="\n".join(page_texts), tables=tuple(tables))


def read_xlsx_content(data: bytes) -> DocumentContent:
    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        tables = tuple(_read_worksheet(sheet) for sheet in workbook.worksheets)
    finally:
        workbook.close()
    text = "\n".join(
        "\t".join(cell.text for cell in row if cell.text) for table in tables for row in table.rows
    )
    return DocumentContent(text=text, tables=tables)


def read_csv_content(data: bytes) -> DocumentContent:
    decoded = _decode_text(data)
    dialect_delimiter = _sniff_delimiter(decoded[:_CSV_SAMPLE_BYTES])
    rows = list(csv.reader(io.StringIO(decoded), delimiter=dialect_delimiter))
    return DocumentContent(text=decoded, tables=(build_raw_table(rows),))


def build_raw_table(
    rows: list[list[str | None]], page: int | None = None, sheet: str | None = None
) -> RawTable:
    return RawTable(
        rows=tuple(
            tuple(
                TableCell(text=_clean_cell_text(cell), reference=f"L{row_number}C{column_number}")
                for column_number, cell in enumerate(row, start=1)
            )
            for row_number, row in enumerate(rows, start=1)
        ),
        page=page,
        sheet=sheet,
    )


def _read_worksheet(sheet) -> RawTable:
    rows = []
    for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        rows.append(
            tuple(
                _build_xlsx_cell(value, f"{get_column_letter(column_number)}{row_number}")
                for column_number, value in enumerate(values, start=1)
            )
        )
    return RawTable(rows=tuple(rows), sheet=sheet.title)


def _build_xlsx_cell(value: object, reference: str) -> TableCell:
    if isinstance(value, bool) or value is None:
        return TableCell(text="" if value is None else str(value), reference=reference)
    if isinstance(value, int | float):
        return TableCell(text=str(value), number=Decimal(str(value)), reference=reference)
    if isinstance(value, datetime | date):
        return TableCell(text=value.strftime("%d/%m/%Y"), reference=reference)
    return TableCell(text=_clean_cell_text(str(value)), reference=reference)


def _clean_cell_text(value: str | None) -> str:
    return " ".join((value or "").split())


def _decode_text(raw: bytes) -> str:
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("latin-1")


def _sniff_delimiter(sample: str) -> str:
    # Brazilian exports use ";" because "," is the decimal separator inside money values.
    for delimiter in _CSV_DELIMITER_PREFERENCE:
        if delimiter in sample:
            return delimiter
    return ","
